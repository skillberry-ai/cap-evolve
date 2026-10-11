#!/usr/bin/env python3
"""Build the conservative SpreadsheetBench-912 residual pool manifest. Zero LLM. Does NOT sample."""
import csv, hashlib, json, os, re, sys, types, importlib.util
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
REPO = Path("/home/dev/projects2/cap-evolve-campaign"); B = Path("/home/dev/projects2/campaign-scratch/bench")
CSV = REPO / "research/capevolve-paper/audit/sb912_residual_exposure.csv"
D912, D400, D200 = (B / "data/extracted" / x for x in ("all_data_912_v0.1", "spreadsheetbench_verified_400", "sample_data_200"))
pop258 = json.load(open(B / "out/pop/AB_912_lo258.json")); pop71 = json.load(open(B / "out/pop/AB_912_lo71.json"))
stub = types.ModuleType("cap_evolve")
for n in ("CapabilityAdapter", "Rollout", "Score", "Task"): setattr(stub, n, type(n, (), {}))
sys.modules["cap_evolve"] = stub; sys.path[:0] = [str(REPO / "templates/adapters/spreadsheetbench"), str(REPO / "templates/adapters")]
os.environ.update(SPREADSHEETBENCH_HARNESS_DIR=str(REPO / "third_party/spreadsheetbench"), SPREADSHEETBENCH_DATA_DIR=str(D912))
spec = importlib.util.spec_from_file_location("ad", REPO / "templates/adapters/spreadsheetbench/adapter.py"); ad = importlib.util.module_from_spec(spec); spec.loader.exec_module(ad)
import openpyxl
def dig(ids): return hashlib.sha256("\n".join(sorted(map(str, ids))).encode()).hexdigest()
rows = list(csv.DictReader(open(CSV)))
ds = {str(e["id"]): e for e in json.load(open(D912 / "dataset.json"))}
ids400 = {str(e["id"]) for e in json.load(open(D400 / "dataset.json"))}
ds400 = json.load(open(D400 / "dataset.json")); ds200 = {str(e["id"]) for e in json.load(open(D200 / "dataset.json"))}
def toks(s): return set(re.findall(r"[a-z0-9]+", s.lower()))
t400 = [(str(e["id"]), toks(e["instruction"])) for e in ds400]
def jacc(a, b): return len(a & b) / len(a | b) if a | b else 0.0
fam = lambda i: i.split("-")[0] if re.fullmatch(r"\d+-\d+", i) else ("CF" if i.startswith("CF_") else i)
dev_exposed = {r["id"] for r in rows if r["top_role"] == "optimization"} | ids400 | ds200
fam_members = Counter(fam(i) for i in ds)
fam_dev = Counter(fam(i) for i in dev_exposed if i in ds)
pa = {r["task"]: r for r in pop258["results"]}; pb = {r["task"]: r for r in pop71["results"]}
def okA(t): return all(c["A"] for c in pa[t]["cases"])
def okB258(t): return all(c.get("B") for c in pa[t]["cases"])
def okB71(t): return all(c.get("B") for c in pb[t]["cases"])
excluded = Counter(); pool = []; excl_list = []
KNOWN_DEFECT = {"43026", "46444", "4714", "52964"}
for r in rows:
    i = r["id"]; why = None
    if r["top_role"] not in ("finalize-score-only", "train-assigned-only"): why = "dev-exposed (optimization role)"
    elif r["in_sample200"] == "True": why = "in public sample_200"
    elif int(r["smoke_runs"]) or int(r["pilot_val_runs"]) or "pilot" in r["roles"] or "opt-" in r["roles"]: why = "used in smoke/pilot design"
    elif i in KNOWN_DEFECT or len(ds_cases := pa[i]["cases"]) != 3: why = "known defective: != 3 graded cases"
    elif not okA(i): why = "unscorable even with gold as output (official comparator)"
    elif not okB258(i): why = "gold-as-output fails after LibreOffice 25.8.7.3 recalculation"
    if why: excluded[why] += 1; excl_list.append({'id': i, 'reason': why}); continue
    pool.append(r)
def feat(r):
    i = r["id"]; e = ds[i]; d = D912 / "spreadsheet" / i
    g = ad._resolve_case_file(d, 1, i, "answer")
    nf = ncell = 0
    try:
        wb = openpyxl.load_workbook(g, data_only=False)
        for part in e["answer_position"].split(","):
            sh, rg = (part.split("!") + [None])[:2] if "!" in part else (wb.sheetnames[0], part)
            sh = sh.strip("'").strip(); rg = rg.strip("'")
            ws = wb[sh] if sh in wb.sheetnames else wb.worksheets[0]
            for row in (ws[rg] if ":" in rg else ((ws[rg],),)):
                for c in row:
                    ncell += 1; nf += c.data_type == "f"
    except Exception as ex:
        return i, None, None
    return i, ncell, nf
with ThreadPoolExecutor(16) as ex: fe = {i: (n, f) for i, n, f in ex.map(feat, pool)}
man = []
for r in pool:
    i = r["id"]; e = ds[i]; n, f = fe[i]; tk = toks(e["instruction"])
    j400 = max((jacc(tk, t) for k, t in t400), default=0)
    man.append({"id": i, "family": fam(i), "family_size_in_912": fam_members[fam(i)], "family_dev_exposed_members": fam_dev.get(fam(i), 0) - (1 if i in dev_exposed else 0),
                "instruction_type": e["instruction_type"], "n_cases": 3, "exposure": r["top_role"], "full_split": r["full_split"],
                "answer_cells": n, "answer_cells_with_formula_in_gold": f, "gold_has_formula_in_range": bool(f) if f is not None else None,
                "instruction_chars": len(e["instruction"]), "max_jaccard_to_verified400_instruction": round(j400, 3),
                "gold_as_output_ok_A": True, "gold_as_output_ok_B_lo258": True, "gold_as_output_ok_B_lo71": okB71(i), "id_type": type(e["id"]).__name__})
man.sort(key=lambda m: m["id"])
summ = {"n_exposure_csv": len(rows), "n_eligible": len(man), "excluded_by_reason": dict(excluded),
        "by_exposure": dict(Counter(m["exposure"] for m in man)), "by_instruction_type": dict(Counter(m["instruction_type"] for m in man)),
        "by_gold_formula": dict(Counter(str(m["gold_has_formula_in_range"]) for m in man)),
        "family_overlap_with_dev_exposed": sum(1 for m in man if m["family_dev_exposed_members"] > 0),
        "families": len({m["family"] for m in man}), "near_dup_ge_0.8_to_verified400": sum(1 for m in man if m["max_jaccard_to_verified400_instruction"] >= .8),
        "needs_lo258_only_for_B": sum(1 for m in man if not m["gold_as_output_ok_B_lo71"]),
        "sorted_ids_sha256": dig([m["id"] for m in man])}
summ["n_if_recalc_only_failures_kept"] = len(man) + sum(1 for x in excl_list if x["reason"].startswith("gold-as-output fails after LibreOffice"))
summ["sorted_ids_sha256_incl_recalc_only_failures"] = dig([m["id"] for m in man] + [x["id"] for x in excl_list if x["reason"].startswith("gold-as-output fails after LibreOffice")])
json.dump({"summary": summ, "ids": man, "excluded": excl_list}, open(sys.argv[1], "w"), indent=1)
print(json.dumps(summ, indent=1))
