#!/usr/bin/env python3
"""E05 (SpreadsheetBench + harness part), zero LLM. REPO env = checkout under test. Prints E05JSON{...}."""
import importlib.util, json, os, shutil, stat, subprocess, sys, tempfile, types, uuid
from pathlib import Path
REPO = Path(os.environ.get("REPO", "/home/dev/projects2/wt-e05")); BENCH = Path("/home/dev/projects2/campaign-scratch/bench")
HARNESS = REPO / "third_party/spreadsheetbench"
for k in [k for k in os.environ if k.startswith("SPREADSHEETBENCH_")]: del os.environ[k]
os.environ["SPREADSHEETBENCH_HARNESS_DIR"] = str(HARNESS)
sys.path[:0] = [str(REPO / "core"), str(REPO / "templates/adapters/spreadsheetbench"), str(REPO / "templates/adapters")]
import openpyxl
from cap_evolve import Rollout, Task
def load():
    spec = importlib.util.spec_from_file_location("e05sb_" + uuid.uuid4().hex[:6], REPO / "templates/adapters/spreadsheetbench/adapter.py")
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m); return m
R = []
def rec(id, area, fault, expected, observed, status, evidence=""):
    R.append({"id": id, "area": area, "fault": fault, "expected": expected, "observed": observed, "status": status, "evidence": evidence})

def wb(path, value, formula=False):
    w = openpyxl.Workbook(); w.active["A1"] = value; w.save(path)

def dataset(root):
    d = root / "spreadsheet" / "t1"; d.mkdir(parents=True)
    wb(d / "1_t1_init.xlsx", "input"); wb(d / "1_t1_golden.xlsx", 2)
    (root / "dataset.json").write_text(json.dumps([{"id": "t1", "instruction": "i", "spreadsheet_path": "spreadsheet/t1", "instruction_type": "Cell-Level Manipulation", "answer_position": "A1"}]))
    return root

def setup(m, tmp, *, recalc=False, soffice=None, harness=None, view=True, allow=False, libre_timeout=None):
    src = dataset(Path(tmp) / "src")
    m.DATA_DIR = str(src); m.AGENT_VIEW = view; m.ALLOW_GOLD_IN_MOUNT = allow; m.GOLD_DIR = ""; m.REWARD_RECALC = recalc; m.SCORING = "hard"
    m._dataset_cache = None; m._VIEWS = {}; m._MOUNT_CHECKED = set(); m.TASK_IDS = []; m._PROVENANCE = None
    m._libreoffice_bin = (lambda: soffice)
    if harness:
        m.HARNESS_DIR = str(harness); m._Vendor._mod = {}
        for n in ("evaluation", "open_spreadsheet", "code_exec", "prompt_format", "llm_api"): sys.modules.pop(n, None)
        sys.path[:] = [x for x in sys.path if not x.startswith(str(HARNESS))]
    if libre_timeout: m.LIBRE_TIMEOUT = libre_timeout
    return src

def score_output(m, writer, tag="t1_0_e05"):
    data = m._data_dir(); d = data / "outputs" / tag; d.mkdir(parents=True, exist_ok=True)
    out = d / "1_t1_output.xlsx"
    if writer: writer(out)
    s = m.Adapter().score(Task(id="t1"), Rollout(task_id="t1", output="", metadata={"run_tag": tag, "id": "t1"}))
    return s, out

# ---- SB-1 malformed outputs -------------------------------------------------------------
for name, w in {"zero_byte_file": lambda p: p.write_bytes(b""), "text_with_xlsx_name": lambda p: p.write_text("not a workbook"),
                "truncated_zip": lambda p: p.write_bytes(open(next(Path(tempfile.gettempdir()).glob("e05_gold.xlsx")), "rb").read()[:200]) , "missing_output_file": None}.items():
    gold_tmp = Path(tempfile.gettempdir()) / "e05_gold.xlsx"; wb(gold_tmp, 2)
    m = load(); tmp = tempfile.mkdtemp(); setup(m, tmp, soffice=None)
    try:
        s, _ = score_output(m, w)
        errored = bool((s.raw or {}).get("errored"))
        ok = s.reward == 0.0 and not errored
        rec(f"SB-1.{name}", "scorer", f"malformed agent output: {name}", "reward 0.0, scored as a candidate failure (not infra, no exception)",
            f"reward={s.reward}, raw.errored={errored}, feedback={ (s.feedback or '')[:70]!r}", "PASS" if ok else "FAIL")
    except Exception as e:
        rec(f"SB-1.{name}", "scorer", f"malformed output {name}", "no exception", f"raised {type(e).__name__}: {e}", "FAIL")

# ---- SB-2 recalculation failure ----------------------------------------------------------
def fake_soffice(dirp, body):
    p = Path(dirp) / "soffice"; p.write_text("#!/usr/bin/env bash\n" + body); p.chmod(0o755); return str(p)
fakes = {
 "exit_nonzero": "exit 1\n",
 "hang_until_timeout": "sleep 20\n",
 "exit0_but_no_file": "exit 0\n",
 "exit0_garbage_file": 'out=""; while [ $# -gt 0 ]; do [ "$1" = "--outdir" ] && out="$2"; last="$1"; shift; done; echo junk > "$out/$(basename "${last%.*}").xlsx"; exit 0\n',
}
formula_out = lambda p: openpyxl.Workbook().save(p) or None
def write_formula(p):
    w = openpyxl.Workbook(); w.active["A1"] = "=1+1"; w.save(p)   # correct as a formula (gold = 2) but no cached value
for fname, body in fakes.items():
    for recalc in (False, True):
        m = load(); tmp = tempfile.mkdtemp(); so = fake_soffice(tmp, body); setup(m, tmp, recalc=recalc, soffice=so, libre_timeout=2)
        s, out = score_output(m, write_formula)
        metrics = {x["name"]: x["value"] for x in s.metrics}
        fb = (s.feedback or "").lower()
        surfaced = "recalc" in fb or "libreoffice" in fb
        raw_flag = any("recalc" in str(k).lower() and "fail" in str(k).lower() for k in (s.raw or {}))
        status = None
        rec(f"SB-2.{fname}.{'recalc_reward' if recalc else 'default_reward'}", "scorer", f"LibreOffice recalculation fails: {fname}; REWARD_RECALC={int(recalc)}",
            "failure is surfaced (metric flag or feedback) and the score is not silently presented as a recalculated one",
            f"reward={s.reward}, hard_restriction={metrics.get('hard_restriction')}, hard_no_recalc={metrics.get('hard_no_recalc')}, failure_in_feedback={surfaced}, failure_flag_in_raw={raw_flag}",
            "PASS" if (surfaced or raw_flag) else ("PARTIAL" if not recalc else "GAP"),
            ("reward is the as-saved value, so unaffected; but the hard_restriction metric is silently the un-recalculated value and nothing marks it" if not recalc else "recalculation is the reward and its failure is not surfaced") if not (surfaced or raw_flag) else "")

# ---- SB-3 gold-named file in the agent-mounted dir ----------------------------------------
m = load(); tmp = tempfile.mkdtemp(); src = setup(m, tmp)
try:
    m.AGENT_VIEW = False; m.ALLOW_GOLD_IN_MOUNT = False; m._data_dir(); r = "no error"
except RuntimeError as e: r = "RuntimeError: " + str(e)[:90]
rec("SB-3.shared_mount_without_allow", "isolation", "historical shared mount (gold beside inputs) without explicit allow", "refused at startup", r, "PASS" if r.startswith("RuntimeError") else "FAIL")
m = load(); tmp = tempfile.mkdtemp(); src = setup(m, tmp); view = m._data_dir(); (view / "spreadsheet/t1/1_t1_answer.xlsx").write_bytes(b"x")
m._MOUNT_CHECKED = set()
try: m._data_dir(); r = "no error"
except RuntimeError as e: r = "RuntimeError"
rec("SB-3.gold_file_in_prebuilt_view_at_startup", "isolation", "gold-named file present in the view when the run starts", "refused at startup", r, "PASS" if r == "RuntimeError" else "FAIL")
m = load(); tmp = tempfile.mkdtemp(); src = setup(m, tmp); view = m._data_dir()   # check passes and is cached
(view / "spreadsheet/t1/1_t1_answer.xlsx").write_bytes(b"x")
try: m._data_dir(); r = "no error"
except RuntimeError as e: r = "RuntimeError"
rec("SB-3.gold_file_appears_after_startup", "isolation", "gold-named file appears in the view AFTER the startup check (mid-run)", "detected before the next evaluation", r + (" (check is cached per mount for the process lifetime)" if r == "no error" else ""),
    "PASS" if r == "RuntimeError" else "GAP", "startup-only assertion; a mid-run write into the mounted dir by agent code is not re-scanned (agent code can write inside /mnt/data)")
m = load(); tmp = tempfile.mkdtemp(); src = setup(m, tmp); v = m._build_agent_view(src, Path(tmp) / "v"); (v / "spreadsheet/t1/1_t1_solution.xlsx").write_bytes(b"x"); m.DATA_DIR = str(v); m.GOLD_DIR = str(src); m._MOUNT_CHECKED = set()
try: m._data_dir(); r = "no error"
except RuntimeError: r = "RuntimeError"
rec("SB-3.unrecognised_filename_in_view", "isolation", "file with a non-input, non-gold name (e.g. renamed gold) in the view", "refused (allowlist)", r, "PASS" if r == "RuntimeError" else "FAIL")

# ---- SB-4 mutated scorer file -------------------------------------------------------------
mut = Path(tempfile.mkdtemp()) / "harness"; shutil.copytree(HARNESS, mut)
ev = mut / "evaluation/evaluation.py"; txt = ev.read_text()
ev.write_text(txt.replace("def compare_workbooks(gt_file, proc_file, instruction_type, answer_position):", "def compare_workbooks(gt_file, proc_file, instruction_type, answer_position):\n    return True, ''\n\ndef _orig_compare_workbooks(gt_file, proc_file, instruction_type, answer_position):", 1))
pinned = subprocess.run(["sha256sum", str(HARNESS / "evaluation/evaluation.py")], capture_output=True, text=True).stdout.split()[0]
m = load(); tmp = tempfile.mkdtemp(); setup(m, tmp, harness=mut)
s, _ = score_output(m, lambda p: p.write_bytes(b"garbage"))
prov = m._scorer_provenance()
rec("SB-4.mutated_scorer_effect", "integrity", "scorer evaluation.py mutated to always return True", "(shows the attack matters) malformed output scores 1.0 under the mutated scorer", f"reward={s.reward}", "INFO" if s.reward == 1.0 else "FAIL")
rec("SB-4.mutation_recorded_in_score", "integrity", "same mutation", "evaluation.py sha256 in score.raw['scorer'] differs from the pinned hash, so the run is attributable after the fact",
    f"recorded={str(prov.get('evaluation_py_sha256'))[:12]} pinned={pinned[:12]} differs={prov.get('evaluation_py_sha256') != pinned}", "PARTIAL" if prov.get("evaluation_py_sha256") != pinned else "FAIL",
    "detection is post hoc: the adapter records the hash but nothing compares it to a pinned value or refuses to run")
frozen = Path(tempfile.mkdtemp()) / "tree"; shutil.copytree(HARNESS, frozen)
man = Path(tempfile.mkdtemp()) / "MANIFEST.json"
subprocess.run([sys.executable, str(BENCH / "isolation/manifest.py"), "write", str(frozen), str(man)], capture_output=True, check=True)
(frozen / "evaluation/evaluation.py").write_text(ev.read_text())
p = subprocess.run([sys.executable, str(BENCH / "isolation/manifest.py"), "verify", str(frozen), str(man)], capture_output=True, text=True)
rec("SB-4.manifest_verify_catches_mutation", "integrity", "mutated evaluation.py in a manifest-protected tree", "verification fails before evaluation", f"exit={p.returncode}: {p.stdout.strip().splitlines()[-1][:80] if p.stdout.strip() else ''}",
    "PASS" if p.returncode != 0 else "FAIL", "mechanism is the prototype in campaign-scratch/bench/isolation (manifest.py + verify_then_run.sh); it is NOT yet wired into run_suite.sh / cap-evolve run")

# ---- SB-5 errored trial: candidate vs infra -------------------------------------------------
m = load(); tmp = tempfile.mkdtemp(); setup(m, tmp, soffice=None)
cand = Path(tmp) / "cand"; shutil.copytree(REPO / "templates/adapters/spreadsheetbench/seed_capability", cand)
(cand / "task_template.md").write_text("no placeholders at all\n")
r = m.Adapter().run_target(Task(id="t1"), cand, seed=0)
rec("SB-5.broken_task_template", "error_tagging", "candidate edits task_template.md and drops required placeholders", "trial errors immediately (no LLM call) and is tagged error_kind=candidate",
    f"error={str(r.error)[:60]!r}, error_kind={(r.metadata or {}).get('error_kind')}", "PASS" if (r.metadata or {}).get("error_kind") == "candidate" else "FAIL")
s = m.Adapter().score(Task(id="t1"), Rollout(task_id="t1", error="LLM call failed: Connection error", metadata={"error_kind": "infra"}))
rec("SB-5.infra_error_scoring", "error_tagging", "gateway connection error", "reward 0.0 with infra wording in feedback; harness drops it as missing data", f"reward={s.reward}, feedback={(s.feedback or '')[:60]!r}", "PASS" if s.reward == 0.0 and "nfrastructure" in (s.feedback or "") else "FAIL")

# harness-level: how drop|fail treat candidate- vs infra-tagged errored trials
sys.path.insert(0, str(REPO / "core/tests"))
import test_paper_mode_switches as T
from cap_evolve import harness
res = {}
for pol in ("drop", "fail"):
    tmpd = Path(tempfile.mkdtemp())
    adapter, run_dir, base = T._setup(tmpd, {"pass": T.TASKS})
    cdir = tmpd / "cand"; shutil.copytree(run_dir.candidate_dir("seed"), cdir)
    (cdir / "cfg.json").write_text(json.dumps({"pass": T.TASKS, "err_candidate": T.TASKS[:6], "err_infra": T.TASKS[6:8]}))
    out = harness.evaluate_candidate(adapter, cdir, run_dir=run_dir, split="val", tag="e05_" + pol, errored_policy=pol)
    pt = {p["task_id"]: p for p in out.per_task} if hasattr(out, "per_task") else {}
    res[pol] = {"mean_reward": round(out.reward, 4), "n_tasks": len(T.TASKS), "errored_tasks": sum(1 for p in out.per_task if (p.get("raw") or {}).get("errored")) if hasattr(out, "per_task") else None}
rec("SB-5.harness_errored_policy", "error_tagging", "candidate crashes on 6/30 tasks (candidate-tagged), gateway errors on 2/30 (infra)", "policy=fail counts the 6 candidate-tagged crashes as 0.0 and still drops the 2 infra errors; policy=drop removes all 8",
    json.dumps(res), "PASS" if res["fail"]["mean_reward"] < res["drop"]["mean_reward"] else "FAIL",
    "default is `drop`: a candidate that crashes on hard tasks looks BETTER under drop; paper mode must set errored_trial_policy: fail")
print("E05JSON" + json.dumps(R))
